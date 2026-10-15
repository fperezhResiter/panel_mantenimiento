"""Verifica independencia del Excel, equivalencia de cálculos y carga atómica."""
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from openpyxl import Workbook
from unittest.mock import patch
from app.base_datos import actualizar_base_datos, leer_fuentes, guardar_configuracion, leer_configuracion
from app.excel import TIPO_ADC
from app.reportes import crear_reporte
from app.seguimiento_hoja import crear_seguimiento
from app.programa import crear_programa


class BaseDatosTest(unittest.TestCase):
    def guardar_gdm(self, zona, nombre_linea='RESIDUOS'):
        ruta = Path(self.temp.name) / 'zonas.xlsx'
        libro = Workbook()
        libro.active.title = 'Prog Resiter'
        libro.active.append([None,None,'Unidad','Centro de Costos','LN','Línea de Negocio','Zona'])
        libro.active.append([None,None,'CECO A',1001,1,nombre_linea,zona])
        libro.save(ruta)
        libro.close()
        return ruta

    def test_actualizar_importa_cambios_de_zona_y_linea(self):
        import base64, gzip, json
        from app.zonas import ClasificadorZonas
        catalogo = Path(self.temp.name) / 'unidades_zonas.json'
        publicacion = Path(self.temp.name) / 'publicacion.js'
        for zona, linea in [(3,'RESIDUOS'), (6,'SANITARIOS')]:
            gdm = self.guardar_gdm(zona, linea)
            actualizar_base_datos(self.excel, self.excel, self.bd, publicacion, gdm, catalogo)
            _, maestro, _ = leer_fuentes(self.bd)
            self.assertEqual(ClasificadorZonas(maestro).resolver('CECO A','ABCD12')['zona'], f'ZONA {zona}')
            self.assertEqual(ClasificadorZonas(maestro).resolver('CECO A','ABCD12')['linea'], linea)
            self.assertEqual(json.loads(catalogo.read_text(encoding='utf-8'))['unidades'][0]['zona'], f'ZONA {zona}')
            p = json.loads(gzip.decompress(base64.b64decode(publicacion.read_text(encoding='utf-8').splitlines()[-1].split(chr(34))[1])))
            self.assertEqual(p['reportes']['2026-09-22']['detalle_zona'][0]['region'], f'ZONA {zona}')

    def test_error_de_publicacion_restaura_bd_catalogo_y_publicacion(self):
        catalogo = Path(self.temp.name) / 'unidades_zonas.json'
        publicacion = Path(self.temp.name) / 'publicacion.js'
        gdm = self.guardar_gdm(3)
        actualizar_base_datos(self.excel, self.excel, self.bd, publicacion, gdm, catalogo)
        archivos = [self.bd, catalogo, publicacion]
        anteriores = [p.read_bytes() for p in archivos]
        self.guardar_gdm(6)
        import os
        reemplazar = os.replace
        def fallo(origen, destino):
            if Path(destino) == publicacion: raise PermissionError('Publicación bloqueada')
            return reemplazar(origen, destino)
        with patch('app.base_datos.os.replace', side_effect=fallo):
            with self.assertRaises(PermissionError):
                actualizar_base_datos(self.excel, self.excel, self.bd, publicacion, gdm, catalogo)
        self.assertEqual([p.read_bytes() for p in archivos], anteriores)
        with self.assertRaises(FileNotFoundError):
            actualizar_base_datos(self.excel, self.excel, self.bd, publicacion, gdm.with_name('ausente.xlsx'), catalogo)
        self.assertEqual([p.read_bytes() for p in archivos], anteriores)

    def test_conserva_configuracion_y_registros(self):
        import sqlite3
        from contextlib import closing
        actualizar_base_datos(self.excel, self.excel, self.bd)
        config = dict(regiones=['REGION BHP'], cecos=[], catalogo=[], centros=[], ocultos=['resumen.tabla'], calendario_excluidas=['ABCD12'])
        guardar_configuracion(self.bd, config)
        with closing(sqlite3.connect(self.bd)) as c, c:
            c.execute('CREATE TABLE registros (id INTEGER PRIMARY KEY, comentario TEXT)')
            c.execute("INSERT INTO registros VALUES (1, 'Conservar')")
        actualizar_base_datos(self.excel, self.excel, self.bd)
        self.assertEqual(leer_configuracion(self.bd), config)
        with closing(sqlite3.connect(self.bd)) as c:
            self.assertEqual(c.execute('SELECT comentario FROM registros').fetchone()[0], 'Conservar')

    def test_configuracion_actualiza_publicacion_y_conserva_datos(self):
        import base64, gzip, json
        from app.publicacion import publicar_consulta
        actualizar_base_datos(self.excel, self.excel, self.bd)
        destino = Path(self.temp.name) / 'publicacion.js'
        publicar_consulta(self.bd, destino)
        def leer():
            comprimido = destino.read_text(encoding='utf-8').split('window.PublicacionMantenimientoGzip = "')[1].split('"')[0]
            return json.loads(gzip.decompress(base64.b64decode(comprimido)))
        anterior = leer()
        config = dict(regiones=['REGION BHP'], cecos=['CECO A'], catalogo=[], centros=[], ocultos=[],
                      zonas=['ZONA NORTE'], catalogo_zonas=['ZONA NORTE'],
                      unidades_regiones=['["REGION BHP","CECO A"]'], unidades_zonas=['["ZONA NORTE","CECO A"]'])
        guardar_configuracion(self.bd, config, destino)
        actual = leer()
        self.assertEqual(actual.pop('configuracion'), config)
        anterior.pop('configuracion')
        self.assertEqual(actual, anterior)
        self.assertEqual(leer_configuracion(self.bd), config)
        with patch('app.publicacion.os.replace', side_effect=PermissionError('Bloqueado')):
            with self.assertRaises(PermissionError):
                guardar_configuracion(self.bd, {**config, 'zonas': []}, destino)
        self.assertEqual(leer_configuracion(self.bd), config)
        self.assertEqual(leer()['configuracion'], config)

    def test_configuracion_crea_publicacion_ausente(self):
        actualizar_base_datos(self.excel, self.excel, self.bd)
        config = dict(regiones=[], cecos=[], catalogo=[], centros=[], ocultos=[])
        destino = Path(self.temp.name) / 'publicacion.js'
        guardar_configuracion(self.bd, config, destino)
        self.assertTrue(destino.is_file())

    def test_publicacion_equivalente_y_error_conserva_bd(self):
        import base64, gzip, json
        from app.publicacion import crear_publicacion
        actualizar_base_datos(self.excel, self.excel, self.bd)
        config = dict(regiones=[], cecos=[], catalogo=[], centros=[], ocultos=[], calendario_excluidas=['ABCD12'])
        guardar_configuracion(self.bd, config)
        contenido = crear_publicacion(self.bd).decode().split('window.PublicacionMantenimientoGzip = "')[1].split('"')[0]
        p = json.loads(gzip.decompress(base64.b64decode(contenido)))
        self.assertEqual(p['configuracion']['calendario_excluidas'], ['ABCD12'])
        excel, maestro, _ = leer_fuentes(self.bd)
        esperado = crear_seguimiento(excel, maestro, date(2026,9,22), date(2026,9,1), incluir_estado=True)
        from app.zonas import enriquecer_zonas
        enriquecer_zonas(esperado, excel, maestro)
        actual = p['intervalos']['0:1']; actual.pop('bd_actualizada')
        self.assertEqual(actual, esperado)
        anterior = self.bd.read_bytes()
        with patch('app.publicacion.publicar_consulta', side_effect=OSError('Sin permiso')):
            with self.assertRaises(OSError):
                actualizar_base_datos(self.excel, self.excel, self.bd, Path(self.temp.name)/'publicacion.js')
        self.assertEqual(self.bd.read_bytes(), anterior)

    def test_error_al_publicar_restaura_bd_y_consulta(self):
        import os
        actualizar_base_datos(self.excel, self.excel, self.bd)
        anterior = self.bd.read_bytes()
        publicacion = Path(self.temp.name)/'publicacion.js'
        publicacion.write_text('Consulta anterior', encoding='utf-8')
        reemplazar = os.replace
        def fallo(origen, destino):
            if Path(destino) == publicacion:
                raise PermissionError('Publicación bloqueada')
            reemplazar(origen, destino)
        with patch('app.base_datos.os.replace', side_effect=fallo), self.assertRaises(PermissionError):
            actualizar_base_datos(self.excel, self.excel, self.bd, publicacion)
        self.assertEqual(self.bd.read_bytes(), anterior)
        self.assertEqual(publicacion.read_text(encoding='utf-8'), 'Consulta anterior')

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.excel = Path(self.temp.name) / 'control.xlsx'
        self.bd = Path(self.temp.name) / 'mantenimiento.sqlite3'
        libro = Workbook()
        libro.active.title = 'Sheet1'
        libro.active.append(['Hora de finalización', 'CeCo (Centro de costo)', 'REGION', 'PATENTE', 'Reportabilidad a realizar'])
        libro.active.append([datetime(2026, 9, 22, 10), 'CECO A', '', 'ABCD12', TIPO_ADC])
        hoja = libro.create_sheet('BD ACTIVOS MOVILES')
        hoja.append(['PATENTE', 'NOMBRE CeCo ACTUAL', 'REGION ACTUAL'])
        hoja.append(['ABCD12', 'CECO A', 'REGION BHP'])
        hoja = libro.create_sheet('Seguimiento KM-HR')
        hoja.append(['Título'])
        hoja.append(['PATENTE', 'REGION ACTUAL', None, 'KM U HR UM', 'UN UM', 'FECHA UM',
                     'INTERVALO', 'UN IN', 'STATUS EQUIPO', date(2026, 9, 1), date(2026, 9, 22)])
        hoja.append(['ABCD12', 'REGION BHP', 'CECO A', 1000, 'Kms', date(2026, 8, 1), 500, 'Kms',
                     'Operativo', 1100, 1400])
        hoja = libro.create_sheet('PROGRAMA')
        hoja.append(['REGION ACTUAL', 'NOMBRE CeCo ACTUAL', 'PATENTE', 'FECHA PROYECTADA', 'ESTADO'])
        hoja.append(['REGION BHP', 'CECO A', 'ABCD12', date(2026, 9, 25), 'REALIZADA'])
        libro.save(self.excel)
        libro.close()

    def test_calculos_identicos_sin_excel(self):
        funciones = [lambda e, m: crear_reporte(e, m, date(2026, 9, 22)),
                     lambda e, m: crear_seguimiento(e, m, date(2026, 9, 22)),
                     lambda e, m: crear_programa(e, '2026-09')]
        esperados = [f(self.excel, self.excel) for f in funciones]
        actualizar_base_datos(self.excel, self.excel, self.bd)
        self.excel.unlink()
        with patch('app.excel.load_workbook', side_effect=AssertionError('No abrir Excel')):
            excel, maestro, estado = leer_fuentes(self.bd)
            self.assertTrue(estado['actualizado'])
            for funcion, esperado in zip(funciones, esperados):
                actual = funcion(excel, maestro)
                actual.pop('actualizado', None)
                esperado.pop('actualizado', None)
                self.assertEqual(actual, esperado)

    def test_carga_fallida_conserva_bd(self):
        actualizar_base_datos(self.excel, self.excel, self.bd)
        anterior = self.bd.read_bytes()
        with self.assertRaises(FileNotFoundError):
            actualizar_base_datos(self.excel, Path(self.temp.name) / 'no-existe.xlsx', self.bd)
        self.assertEqual(anterior, self.bd.read_bytes())
        self.assertEqual(list(self.bd.parent.glob('mantenimiento-*.sqlite3')), [])

    def test_recarga_sustituye_datos(self):
        actualizar_base_datos(self.excel, self.excel, self.bd)
        from openpyxl import load_workbook
        libro = load_workbook(self.excel)
        libro['Sheet1'].delete_rows(2)
        libro.save(self.excel)
        libro.close()
        actualizar_base_datos(self.excel, self.excel, self.bd)
        excel, maestro, _ = leer_fuentes(self.bd)
        self.assertEqual(crear_reporte(excel, maestro, date(2026, 9, 22))['reportadas'], 0)
