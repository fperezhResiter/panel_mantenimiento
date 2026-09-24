



"""Extract typed Excel data using ClientContext; see connection_sharepoint.md."""

import argparse
import getpass
import json
import math
import os
from dataclasses import dataclass
from datetime import date, datetime, time, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, Any, Callable
from urllib.parse import urlsplit

if TYPE_CHECKING:
    from office365.sharepoint.client_context import ClientContext


APP_DIRECTORY = Path(__file__).resolve().parent
SUPPORTED_TYPES = ("STRING", "INT64", "FLOAT64", "NUMERIC", "BIGNUMERIC",
                   "BOOL", "DATE", "DATETIME", "TIME", "TIMESTAMP")
TYPE_ALIASES = {"INTEGER": "INT64", "INT": "INT64", "FLOAT": "FLOAT64",
                "BOOLEAN": "BOOL", "DECIMAL": "NUMERIC", "BIGDECIMAL": "BIGNUMERIC"}
NUMERIC_MAX = Decimal("99999999999999999999999999999.999999999")
BIGNUMERIC_MAX = Decimal(
    "578960446186580977117854925043439539266.34992332820282019728792003956564819967"
)
BIGNUMERIC_MIN = Decimal(
    "-578960446186580977117854925043439539266.34992332820282019728792003956564819968"
)

DEFAULT_FORMATS = {"DATE": "%Y-%m-%d", "DATETIME": "%Y-%m-%d %H:%M:%S",
                   "TIME": "%H:%M:%S", "TIMESTAMP": "%Y-%m-%d %H:%M:%S%z"}


class ConnectionSharePointError(Exception):
    """An actionable connection, configuration or extraction failure."""


@dataclass(frozen=True)
class SharePointSettings:
    site_url: str

    def __post_init__(self) -> None:
        if not self.site_url or "TU_" in self.site_url:
            raise ValueError("Completa SHAREPOINT_SITE_URL en el .env junto a connection_sharepoint.py.")
        site = urlsplit(self.site_url)
        if (site.scheme != "https" or not site.hostname
                or not site.hostname.endswith(".sharepoint.com")
                or site.query or site.fragment or site.username or site.port):
            raise ValueError("SHAREPOINT_SITE_URL debe ser la URL HTTPS del sitio SharePoint Online.")


@dataclass(frozen=True)
class ColumnSchema:
    name: str
    data_type: str
    nullable: bool = True
    date_format: str | None = None
    decimal_separator: str = "."
    thousands_separator: str = ""

    def __post_init__(self) -> None:
        canonical_type = self.data_type.strip().upper()
        canonical_type = TYPE_ALIASES.get(canonical_type, canonical_type)
        object.__setattr__(self, "data_type", canonical_type)
        if not self.name.strip() or canonical_type not in SUPPORTED_TYPES:
            raise ValueError(f"Columna o tipo inválido: {self.name!r}, {canonical_type}.")
        if self.decimal_separator not in {".", ","}:
            raise ValueError("El separador decimal debe ser punto o coma.")
        if self.thousands_separator not in {"", ".", ",", " "}:
            raise ValueError("Separador de miles inválido.")
        if self.decimal_separator == self.thousands_separator:
            raise ValueError("Los separadores decimal y de miles deben ser distintos.")


@dataclass(frozen=True)
class ExtractionRequest:
    library: str
    folder: str
    file_name: str
    sheet_name: str
    columns: tuple[ColumnSchema, ...]
    header_row: int = 1

    def __post_init__(self) -> None:
        if not self.library.strip() or not self.sheet_name.strip():
            raise ValueError("La biblioteca y la hoja son obligatorias.")
        if Path(self.file_name).suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError("El archivo debe tener extensión .xlsx o .xlsm.")
        if "/" in self.file_name or "\\" in self.file_name:
            raise ValueError("Indica las carpetas por separado del nombre del archivo.")
        folder = self.folder.replace("\\", "/").strip("/")
        if any(part in {".", ".."} for part in folder.split("/")):
            raise ValueError("La ruta no puede contener . ni ..")
        object.__setattr__(self, "folder", folder)
        if type(self.header_row) is not int or self.header_row < 1:
            raise ValueError("La fila de encabezados debe ser un entero positivo.")
        names = [column.name for column in self.columns]
        if not names or len(names) != len(set(names)):
            raise ValueError("Selecciona al menos una columna, sin repetir nombres.")


def load_settings(env_path: Path = APP_DIRECTORY / ".env") -> SharePointSettings:
    from dotenv import dotenv_values

    file_values = dotenv_values(env_path, encoding="utf-8-sig")
    site_url = os.environ.get("SHAREPOINT_SITE_URL", file_values.get("SHAREPOINT_SITE_URL"))
    return SharePointSettings(str(site_url or "").strip())


def create_client_context(settings: SharePointSettings, username: str, password: str) -> "ClientContext":
    """Build the requested legacy user context; Online retired this authentication."""
    from office365.runtime.auth.user_credential import UserCredential
    from office365.sharepoint.client_context import ClientContext

    if not username.strip() or not password:
        raise ValueError("El usuario y la contraseña son obligatorios.")
    try:
        return ClientContext(settings.site_url).with_credentials(UserCredential(username.strip(), password))
    except Exception:
        # Do not expose provider errors that may contain authentication data.
        raise ConnectionSharePointError(
            "No se pudo configurar ClientContext con usuario y contraseña. "
            "SharePoint Online retiró este flujo antiguo; consulta connection_sharepoint.md."
        ) from None


class SharePointClient:
    """Resolve a document library and download an Excel without modifying it."""

    def __init__(self, context: "ClientContext") -> None:
        self._context = context

    def download_excel(self, request: ExtractionRequest) -> bytes:
        try:
            library = self._context.web.lists.get_by_title(request.library)
            root_folder = library.root_folder.get().execute_query()
            # Resolve the actual library URL: its title may differ from its URL name.
            root_path = root_folder.properties["ServerRelativeUrl"].rstrip("/")
            file_path = "/".join(part for part in (root_path, request.folder, request.file_name) if part)
            with BytesIO() as output:
                remote_file = self._context.web.get_file_by_server_relative_path(file_path)
                remote_file.download(output).execute_query()
                return output.getvalue()
        except Exception:
            raise ConnectionSharePointError(
                "No se pudo leer el Excel con ClientContext. Revisa biblioteca, ruta, archivo y permisos. "
                "El flujo antiguo de usuario/contraseña fue retirado de SharePoint Online "
                "y no admite MFA. Consulta connection_sharepoint.md."
            ) from None


def parse_number(value: Any, column: ColumnSchema) -> Decimal:
    if isinstance(value, (bool, date, time)):
        raise ValueError("se esperaba un número")
    text = str(value).strip()
    if isinstance(value, str):
        if column.thousands_separator:
            text = text.replace(column.thousands_separator, "")
        text = text.replace(column.decimal_separator, ".")
    number = Decimal(text)
    if not number.is_finite():
        raise ValueError("se requiere un número finito")
    return number


def convert_number(value: Any, column: ColumnSchema) -> int | float | Decimal:
    number = parse_number(value, column)
    if column.data_type == "INT64":
        if number != number.to_integral_value() or not -(2**63) <= number <= 2**63 - 1:
            raise ValueError("se requiere un entero dentro del rango INT64")
        return int(number)
    if column.data_type == "FLOAT64":
        result = float(number)
        if not math.isfinite(result):
            raise ValueError("fuera del rango FLOAT64 finito")
        return result
    maximum, scale = (NUMERIC_MAX, 9) if column.data_type == "NUMERIC" else (BIGNUMERIC_MAX, 38)
    minimum = NUMERIC_MAX.copy_negate() if column.data_type == "NUMERIC" else BIGNUMERIC_MIN
    # Inspect digits without Decimal.normalize(), which can round with the active context.
    _, digits, exponent = number.as_tuple()
    digits = list(digits)
    while digits and digits[-1] == 0:
        digits.pop()
        exponent += 1
    fractional_digits = max(0, -exponent) if digits else 0
    if not minimum <= number <= maximum or fractional_digits > scale:
        raise ValueError(f"fuera del rango o escala de {column.data_type} (máximo {scale} decimales)")
    return number


def convert_boolean(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float, Decimal)) and value in (0, 1):
        return bool(value)
    text = str(value).strip().casefold()
    if text in {"true", "verdadero", "si", "sí", "1"}:
        return True
    if text in {"false", "falso", "no", "0"}:
        return False
    raise ValueError("se esperaba true/false, sí/no o 1/0")


def convert_temporal(value: Any, column: ColumnSchema) -> date | datetime | time:
    kind = column.data_type
    if isinstance(value, str):
        parsed = datetime.strptime(value.strip(), column.date_format or DEFAULT_FORMATS[kind])
        value = parsed.time() if kind == "TIME" else parsed
    if kind == "TIME":
        if not isinstance(value, time) or value.tzinfo is not None:
            raise ValueError("TIME requiere una hora sin zona horaria")
        return value
    if isinstance(value, date) and not isinstance(value, datetime):
        value = datetime.combine(value, time.min)
    if not isinstance(value, datetime):
        raise ValueError("se esperaba una fecha Excel o texto con el formato indicado")
    if kind == "TIMESTAMP":
        if value.tzinfo is None or value.utcoffset() is None:
            raise ValueError("TIMESTAMP requiere zona horaria explícita, por ejemplo -0300")
        return value.astimezone(timezone.utc)
    if value.tzinfo is not None:
        raise ValueError(f"{kind} no admite zona horaria")
    if kind == "DATE":
        if value.time() != time.min:
            raise ValueError("DATE no admite una hora distinta de medianoche; usa DATETIME")
        return value.date()
    return value


def convert_value(value: Any, column: ColumnSchema) -> Any:
    if value is None or isinstance(value, str) and not value.strip():
        if not column.nullable:
            raise ValueError("la columna es REQUIRED y el valor está vacío")
        return None
    if column.data_type == "STRING":
        return str(value)
    if column.data_type == "BOOL":
        return convert_boolean(value)
    if column.data_type in DEFAULT_FORMATS:
        return convert_temporal(value, column)
    return convert_number(value, column)


def extract_records(content: bytes, request: ExtractionRequest) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if request.sheet_name not in workbook.sheetnames:
            raise ValueError(f"Hoja no encontrada. Disponibles: {', '.join(workbook.sheetnames)}")
        rows = workbook[request.sheet_name].iter_rows(min_row=request.header_row, values_only=True)
        headers = list(next(rows, ()))
        for column in request.columns:
            if headers.count(column.name) != 1:
                raise ValueError(f"La columna {column.name!r} no existe o está duplicada.")
        indexes = [headers.index(column.name) for column in request.columns]
        records = []
        for row_number, row in enumerate(rows, start=request.header_row + 1):
            if all(value is None for value in row):
                continue
            record = {}
            for column, index in zip(request.columns, indexes):
                try:
                    record[column.name] = convert_value(row[index], column)
                except (ValueError, TypeError, InvalidOperation, OverflowError) as error:
                    raise ValueError(f"Fila {row_number}, columna {column.name!r}: {error}") from error
            records.append(record)
        return records
    finally:
        workbook.close()


def read_sharepoint_excel(
    settings: SharePointSettings, request: ExtractionRequest, username: str, password: str,
) -> list[dict[str, Any]]:
    context = create_client_context(settings, username, password)
    content = SharePointClient(context).download_excel(request)
    return extract_records(content, request)


def prompt_value(label: str, validator: Callable[[str], Any], default: str | None = None) -> Any:
    while True:
        suffix = f" [{default}]" if default is not None else ""
        value = input(f"{label}{suffix}: ").strip()
        try:
            return validator(value if value or default is None else default)
        except ValueError as error:
            print(f"Entrada inválida: {error}")


def required_text(value: str) -> str:
    if not value:
        raise ValueError("este campo es obligatorio")
    return value


def positive_integer(value: str) -> int:
    number = int(value)
    if number < 1:
        raise ValueError("debe ser mayor que cero")
    return number


def yes_or_no(value: str) -> bool:
    if value.casefold() not in {"s", "n"}:
        raise ValueError("escribe s o n")
    return value.casefold() == "s"


def prompt_column() -> ColumnSchema:
    name = prompt_value("Nombre exacto de la columna", required_text)
    data_type = prompt_value("Tipo BigQuery", lambda text: ColumnSchema(name, text).data_type)
    nullable = prompt_value("¿Permite nulos? s/n", yes_or_no, "s")
    date_format = None
    decimal_separator, thousands_separator = ".", ""
    if data_type in DEFAULT_FORMATS:
        date_format = prompt_value("Formato de fechas/horas en texto", required_text, DEFAULT_FORMATS[data_type])
    if data_type in {"INT64", "FLOAT64", "NUMERIC", "BIGNUMERIC"}:
        decimal_separator = prompt_value(
            "Separador decimal en texto", lambda text: ColumnSchema(name, data_type, decimal_separator=text).decimal_separator, "."
        )
        thousands_separator = prompt_value(
            "Separador de miles en texto (Enter = ninguno; ESPACIO = espacio)",
            lambda text: ColumnSchema(name, data_type, decimal_separator=decimal_separator,
                                      thousands_separator=" " if text.upper() == "ESPACIO" else text).thousands_separator,
        )
    return ColumnSchema(name, data_type, nullable, date_format, decimal_separator, thousands_separator)


def prompt_extraction_request() -> ExtractionRequest:
    library = prompt_value("Biblioteca de documentos", required_text)
    folder = input("Ruta dentro de la biblioteca (Enter = raíz): ").strip()
    file_name = prompt_value("Nombre del Excel con extensión", required_text)
    sheet_name = prompt_value("Nombre exacto de la hoja", required_text)
    header_row = prompt_value("Fila de encabezados", positive_integer, "1")
    count = prompt_value("Cantidad de columnas a extraer", positive_integer)
    print("Tipos disponibles: " + ", ".join(SUPPORTED_TYPES))
    columns = []
    while len(columns) < count:
        print(f"\nColumna {len(columns) + 1} de {count}")
        column = prompt_column()
        if any(existing.name == column.name for existing in columns):
            print("La columna ya fue seleccionada; indica otra.")
            continue
        columns.append(column)
    return ExtractionRequest(library, folder, file_name, sheet_name, tuple(columns), header_row)


def serialize_value(value: Any) -> str:
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    raise TypeError(f"Tipo no serializable: {type(value).__name__}")


def save_records(records: list[dict[str, Any]], output_path: Path) -> None:
    payload = json.dumps(records, ensure_ascii=False, indent=2, default=serialize_value, allow_nan=False)
    # Exclusive creation prevents an output path from overwriting .env or source files.
    with output_path.open("x", encoding="utf-8") as output:
        output.write(payload)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.parse_args()
    try:
        settings = load_settings()
        print("Aviso: el flujo UserCredential fue retirado de SharePoint Online; ver connection_sharepoint.md.")
        username = prompt_value("Usuario Microsoft 365 (correo)", required_text)
        password = getpass.getpass("Contraseña (no se muestra ni se guarda): ")
        if not password:
            raise ValueError("La contraseña es obligatoria.")
        request = prompt_extraction_request()
        output_path = Path(prompt_value(
            "Ruta JSON de salida (debe ser un archivo nuevo)", required_text,
            str(APP_DIRECTORY / "connection_sharepoint_output.json"),
        )).expanduser()
        if output_path.exists() or not output_path.parent.is_dir():
            raise ValueError("La salida ya existe o su carpeta no existe; indica un archivo nuevo.")
        records = read_sharepoint_excel(settings, request, username, password)
        save_records(records, output_path)
        print(f"Extraídas {len(records)} filas en {output_path.resolve()}")
    except (EOFError, KeyboardInterrupt):
        parser.exit(1, "\nExtracción cancelada.\n")
    except Exception as error:
        # CLI boundary: library functions retain their original exceptions for callers.
        parser.exit(1, f"Error: {error}\n")


if __name__ == "__main__":
    main()
