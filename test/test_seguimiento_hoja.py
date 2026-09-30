import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from openpyxl import Workbook
from app.seguimiento_hoja import crear_seguimiento
from app.excel import leer_tabla


class SeguimientoHojaTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.ruta = Path(self.temp.name) / 'control.xlsx'
        self.libro = Workbook()
        self.libro.active.title = 'Sheet1'
        maestro = self.libro.create_sheet('BD ACTIVOS MOVILES')
        maestro.append(['PATENTE', 'NOMBRE CeCo ACTUAL', 'REGION ACTUAL'])
        maestro.append(['LBSR70', 'CECO MAESTRO', 'REGION TARAPACA'])
        self.hoja = self.libro.create_sheet('Seguimiento KM-HR')
        self.hoja.append(['Título del reporte'])
        self.hoja.append(['PATENTE', 'REGION ACTUAL', None, 'KM U HR UM', 'UN UM', 'FECHA UM',
                          'INTERVALO', 'UN IN', 'STATUS EQUIPO', '01-09-2026',
                          datetime(2026, 9, 8), '15-09-2026', '22-09-2026', 'FECHA PROYECTADA'])

    def consultar(self, inicio=date(2026, 9, 1), corte=date(2026, 9, 22)):
        self.libro.save(self.ruta)
        return crear_seguimiento(self.ruta, self.ruta, corte, inicio, incluir_estado=True)

    def test_ceros_fechas_y_proyeccion_independiente(self):
        self.hoja.append(['LBSR70', 'REGION TARAPACA', 'SUCURSAL IQUIQUE', 210105, 'Kms',
                          '14-04-2026', 10000, 'Kms', 'Operativo', 220905, 0, 222494, 222875, '01-01-2099'])
        r = self.consultar()
        e = r['equipos'][0]
        self.assertEqual([l['valor'] if l else None for l in e['lecturas']], [220905, None, 222494, 222875])
        self.assertEqual(e['calculo']['saldo'], -2770)
        self.assertEqual(e['calculo']['fecha_proyectada'], '2026-08-24')
        self.assertEqual(e['ceco'], 'SUCURSAL IQUIQUE')
        self.assertEqual(e['estado_equipo'], 'OPERATIVO')
        self.assertIsNone(e['fecha_estado_equipo'])
        self.assertTrue(r['incluye_estado_equipo'])
        self.assertEqual(r['semanas'][0]['desde'], '2026-09-01')

    def test_retroceso_y_blancos_no_se_proyectan(self):
        self.hoja.cell(2, 15, 'Comentario Tania')
        self.hoja.append(['LJVF94', 'REGION TARAPACA', None, 1824, 'Hrs', '10-02-2026',
                          400, 'Hrs', 'Fuera de servicio', 2124, '\u00a0', 2045, 0, None,
                          'Pendiente revisión de Tania'])
        e = self.consultar()['equipos'][0]
        self.assertEqual(e['comentario_tania'], 'Pendiente revisión de Tania')
        self.assertEqual(e['calculo']['estado'], 'STAND BY')
        self.assertIsNone(e['calculo']['fecha_proyectada'])
        self.assertIsNone(e['lecturas'][1])

    def test_fecha_corte_y_duplicados(self):
        fila = ['LBSR70', 'REGION TARAPACA', None, 1000, 'Kms', '01-08-2026',
                500, 'Kms', 'Operativo', 1100, 1240, 1300, 1400, None]
        self.hoja.append(fila)
        self.hoja.append(fila)
        r = self.consultar(corte=date(2026, 9, 8))
        self.assertEqual(len(r['equipos']), 1)
        self.assertEqual(len(r['semanas']), 2)
        self.assertEqual(r['equipos'][0]['calculo']['fecha_proyectada'], '2026-09-21')
        self.hoja.append([*fila[:9], 1200, *fila[10:]])
        self.assertIsNone(self.consultar()['equipos'][0]['calculo']['fecha_proyectada'])

    def test_hoja_explicita_y_periodo_sin_columnas(self):
        self.libro.save(self.ruta)
        with self.assertRaisesRegex(ValueError, 'no se encontró la hoja'):
            leer_tabla(self.ruta, ['PATENTE'], nombre_hoja='Inexistente')
        with self.assertRaisesRegex(ValueError, 'no hay columnas con fecha'):
            self.consultar(inicio=date(2026, 10, 1), corte=date(2026, 10, 8))
