"""Adaptación de la matriz Seguimiento KM-HR al cálculo del portal."""
from datetime import date
from openpyxl import load_workbook
from .excel import normalizar, convertir_fecha, patente_valida, leer_tabla
from .seguimiento import numero, unidad_control, calcular_equipo


def crear_seguimiento(excel, maestro, corte, inicio=date(2026, 8, 4), incluir_estado=False):
    if corte < inicio:
        raise ValueError('La última semana debe ser igual o posterior al inicio de la semana 1.')
    if (corte - inicio).days // 7 + 1 > 104:
        raise ValueError('Selecciona un período de hasta 104 semanas.')
    activos, _, _ = leer_tabla(maestro, ['PATENTE', 'NOMBRE CeCo ACTUAL', 'REGION ACTUAL'],
                              nombre_hoja='BD ACTIVOS MOVILES')
    ubicaciones = {}
    for patente, ceco, region in activos:
        p = patente_valida(patente)
        if p:
            ubicaciones.setdefault(p, set()).add((normalizar(ceco), normalizar(region)))
    libro = load_workbook(excel, read_only=True, data_only=True)
    try:
        hoja = next((h for h in libro if normalizar(h.title) == 'SEGUIMIENTO KM-HR'), None)
        if hoja is None:
            raise ValueError(f'{excel.name}: no se encontró la hoja Seguimiento KM-HR.')
        filas = hoja.iter_rows(values_only=True)
        requeridas = ['PATENTE', 'REGION ACTUAL', 'KM U HR UM', 'UN UM', 'FECHA UM',
                      'INTERVALO', 'UN IN', 'STATUS EQUIPO']
        # Permite títulos sobre la cabecera, sin confundirlos con datos.
        for fila_cabecera, cabecera in enumerate(filas, 1):
            indices = {normalizar(c): i for i, c in enumerate(cabecera) if c is not None}
            if all(c in indices for c in requeridas):
                break
            if fila_cabecera >= 20:
                raise ValueError('Seguimiento KM-HR: faltan columnas requeridas: ' + ', '.join(requeridas))
        else:
            raise ValueError('Seguimiento KM-HR: no se encontró una cabecera válida.')
        fechas = []
        for i, c in enumerate(cabecera):
            dia = convertir_fecha(c, libro.epoch)
            if dia and inicio <= dia <= corte:
                fechas.append((dia, i))
        fechas.sort()
        if not fechas:
            raise ValueError('Seguimiento KM-HR: no hay columnas con fecha dentro del período seleccionado.')
        if len({d for d, _ in fechas}) != len(fechas):
            raise ValueError('Seguimiento KM-HR: hay columnas con fechas duplicadas.')
        semanas = [{'numero': n, 'fecha': d.isoformat(), 'desde': d.isoformat(), 'hasta': d.isoformat()}
                   for n, (d, _) in enumerate(fechas, 1)]
        equipos = {}
        ceros = 0
        for fila_num, fila in enumerate(filas, fila_cabecera + 1):
            def valor(nombre):
                i = indices.get(nombre)
                return fila[i] if i is not None and i < len(fila) else None
            p = patente_valida(valor('PATENTE'))
            if not p:
                continue
            region = normalizar(valor('REGION ACTUAL'))
            ceco = normalizar(valor('NOMBRE CECO ACTUAL'))
            # El encabezado del CeCo actual viene vacío, inmediatamente tras REGION ACTUAL.
            i_ceco = indices['REGION ACTUAL'] + 1
            if not ceco and i_ceco < len(cabecera) and not normalizar(cabecera[i_ceco]):
                ceco = normalizar(fila[i_ceco])
            ubicacion = ubicaciones.get(p, set())
            if len(ubicacion) == 1:
                ceco_maestro, region_maestro = next(iter(ubicacion))
                ceco = ceco or ceco_maestro
                region = region or region_maestro
            unidad = unidad_control(valor('UN UM'))
            fecha_um = convertir_fecha(valor('FECHA UM'), libro.epoch)
            fecha_real = convertir_fecha(valor('FECHA PROYECTADA REAL'), libro.epoch)
            lecturas = []
            for n, (dia, i) in enumerate(fechas, 1):
                bruto = fila[i] if i < len(fila) else None
                lectura = numero(bruto)
                if lectura == 0 or not str(bruto or '').strip():
                    ceros += lectura == 0
                    lecturas.append(None)
                    continue
                lecturas.append({'semana': n, 'valor': lectura, 'unidad': unidad,
                                 'fecha_lectura': dia.isoformat(), 'fecha_envio': None,
                                 'unidad_desde_maestro': False, 'fecha_desde_envio': False,
                                 'fila_excel': fila_num, 'estado_equipo': normalizar(valor('STATUS EQUIPO'))})
            e = {'patente': p, 'region': region or 'SIN REGIÓN', 'ceco': ceco or 'SIN CECO',
                 'ultima_mantencion': numero(valor('KM U HR UM')), 'unidad': unidad,
                 'fecha_mantencion': fecha_um.isoformat() if fecha_um else None,
                 'fecha_proyectada_real': fecha_real.isoformat() if fecha_real else None,
                 'intervalo': numero(valor('INTERVALO')), 'unidad_intervalo': unidad_control(valor('UN IN')),
                 'lecturas': lecturas, 'conflicto_maestro': False,
                 'comentario_tania': str(valor('COMENTARIO TANIA') or '').strip(),
                 'origen_mantencion': hoja.title, 'certificado_mantencion': '',
                 'fila_mantencion': None, 'fecha_envio_mantencion': None,
                 'estado_equipo': normalizar(valor('STATUS EQUIPO')), 'fecha_estado_equipo': None}
            if p in equipos:
                anterior = equipos[p]
                campos = ('region', 'ceco', 'ultima_mantencion', 'unidad', 'fecha_mantencion',
                          'intervalo', 'unidad_intervalo', 'estado_equipo', 'fecha_proyectada_real')
                if any(anterior[k] != e[k] for k in campos) or [l['valor'] if l else None for l in anterior['lecturas']] != [l['valor'] if l else None for l in lecturas]:
                    anterior['conflicto_maestro'] = True
                continue
            equipos[p] = e
        for e in equipos.values():
            e['calculo'] = calcular_equipo(e)
        avisos = ['Lecturas tomadas de las columnas con fecha de Seguimiento KM-HR. '
                  'La última mantención y el estado del equipo reflejan el estado actual de la hoja. '
                  'La fecha proyectada se calcula en el portal; no se usa FECHA PROYECTADA del Excel. '
                  'El calendario prioriza FECHA PROYECTADA REAL cuando contiene una fecha válida.']
        if ceros:
            avisos.append(f'{ceros} celdas con cero se interpretaron como semanas sin reporte.')
        conflictos = sum(e['conflicto_maestro'] for e in equipos.values())
        if conflictos:
            avisos.append(f'{conflictos} patentes repetidas con datos distintos: REVISAR.')
        return {'version': 2, 'fecha': corte.isoformat(), 'inicio': semanas[0]['fecha'],
                'semanas': semanas, 'equipos': sorted(equipos.values(), key=lambda e: (e['region'], e['ceco'], e['patente'])),
                'advertencias': avisos, 'fuente': excel.name, 'maestro': maestro.name, 'hoja': hoja.title,
                'incluye_estado_equipo': incluir_estado, 'actualizaciones_mantencion': 0,
                'limite_mantenciones': corte.isoformat()}
    finally:
        libro.close()
