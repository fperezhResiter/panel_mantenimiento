"""Copia local SQLite de las hojas que consumen los paneles."""
import json
import os
import sqlite3
import tempfile
from contextlib import closing
from datetime import date, datetime, time
from pathlib import Path
from threading import Lock

from openpyxl import load_workbook
from .excel import normalizar

VERSION = '1'
ACTUALIZACION = Lock()
REQUERIDAS = {
    'Sheet1': ['Hora de finalización', 'CeCo (Centro de costo)', 'REGION', 'PATENTE', 'Reportabilidad a realizar'],
    'BD ACTIVOS MOVILES': ['PATENTE', 'NOMBRE CeCo ACTUAL', 'REGION ACTUAL'],
    'Seguimiento KM-HR': ['PATENTE', 'REGION ACTUAL', 'KM U HR UM', 'UN UM', 'FECHA UM', 'INTERVALO', 'UN IN', 'STATUS EQUIPO'],
    'PROGRAMA': ['REGION ACTUAL', 'NOMBRE CeCo ACTUAL', 'PATENTE', 'FECHA PROYECTADA', 'ESTADO'],
}


def codificar(valor):
    if isinstance(valor, (datetime, date, time)):
        return {'tipo': type(valor).__name__, 'valor': valor.isoformat()}
    raise TypeError(f'Tipo de celda no compatible: {type(valor).__name__}')


def decodificar(valor):
    tipos = {'datetime': datetime, 'date': date, 'time': time}
    return tipos[valor['tipo']].fromisoformat(valor['valor'])


class HojaBD:
    def __init__(self, titulo, filas):
        self.title, self.filas = titulo, filas

    def iter_rows(self, values_only=True):
        return iter(self.filas)


class FuenteBD:
    def __init__(self, nombre, epoch, hojas):
        self.name, self.epoch, self.hojas = nombre, epoch, hojas
        self.sheetnames = [h.title for h in hojas]

    def __iter__(self):
        return iter(self.hojas)

    def close(self):
        pass


def actualizar_base_datos(excel, maestro, destino, publicacion=None, zonas=None, catalogo_destino=None):
    """Publica una carga completa; un error conserva intacta la BD anterior."""
    with ACTUALIZACION:
        destino = Path(destino).resolve()
        firmas = {Path(ruta).resolve(): (Path(ruta).stat().st_size, Path(ruta).stat().st_mtime_ns)
                  for ruta in (excel, maestro, *([zonas] if zonas is not None else []))}
        catalogo = None
        catalogo_temporal = None
        if zonas is not None:
            from .catalogo_zonas import leer_catalogo_zonas
            from .zonas import CATALOGO
            catalogo_destino = Path(catalogo_destino or CATALOGO)
            anterior_catalogo = json.loads((catalogo_destino if catalogo_destino.exists() else CATALOGO).read_text(encoding='utf-8'))
            catalogo = leer_catalogo_zonas(zonas, anterior_catalogo)
        destino.parent.mkdir(parents=True, exist_ok=True)
        descriptor, nombre = tempfile.mkstemp(prefix='mantenimiento-', suffix='.sqlite3', dir=destino.parent)
        os.close(descriptor)
        temporal = Path(nombre)
        consulta_temporal = None
        try:
            if destino.exists():
                with closing(conectar(destino)) as anterior, closing(sqlite3.connect(temporal)) as copia:
                    anterior.backup(copia)
            with closing(sqlite3.connect(temporal)) as conexion, conexion:
                conexion.executescript('''
                    DROP TABLE IF EXISTS metadatos;
                    DROP TABLE IF EXISTS fuentes;
                    DROP TABLE IF EXISTS hojas;
                    DROP TABLE IF EXISTS filas;
                    CREATE TABLE metadatos (clave TEXT PRIMARY KEY, valor TEXT NOT NULL);
                    CREATE TABLE fuentes (rol TEXT PRIMARY KEY, nombre TEXT NOT NULL, ruta TEXT NOT NULL, epoch TEXT NOT NULL);
                    CREATE TABLE hojas (rol TEXT NOT NULL, nombre TEXT NOT NULL, PRIMARY KEY (rol, nombre));
                    CREATE TABLE filas (rol TEXT NOT NULL, hoja TEXT NOT NULL, numero INTEGER NOT NULL, datos TEXT NOT NULL,
                                        PRIMARY KEY (rol, hoja, numero));
                ''')
                for rol, ruta, titulos in [('excel', Path(excel), ['Sheet1', 'Seguimiento KM-HR', 'PROGRAMA']),
                                           ('maestro', Path(maestro), ['BD ACTIVOS MOVILES'])]:
                    firma = ruta.stat()
                    libro = load_workbook(ruta, read_only=True, data_only=True)
                    try:
                        conexion.execute('INSERT INTO fuentes VALUES (?, ?, ?, ?)',
                                         (rol, ruta.name, str(ruta.resolve()), libro.epoch.isoformat()))
                        for titulo in titulos:
                            hoja = next((h for h in libro if normalizar(h.title) == normalizar(titulo)), None)
                            if hoja is None:
                                raise ValueError(f'{ruta.name}: falta la hoja {titulo}.')
                            conexion.execute('INSERT INTO hojas VALUES (?, ?)', (rol, hoja.title))
                            valida = False
                            limite = 20 if titulo == 'Seguimiento KM-HR' else 1
                            for numero, fila in enumerate(hoja.iter_rows(values_only=True), 1):
                                if numero <= limite and all(normalizar(c) in {normalizar(v) for v in fila}
                                                           for c in REQUERIDAS[titulo]):
                                    valida = True
                                if numero == limite and not valida:
                                    raise ValueError(f'{ruta.name}: faltan columnas requeridas en {titulo}.')
                                conexion.execute('INSERT INTO filas VALUES (?, ?, ?, ?)',
                                                 (rol, hoja.title, numero, json.dumps(fila, ensure_ascii=False,
                                                                                     default=codificar, allow_nan=False)))
                            if not valida:
                                raise ValueError(f'{ruta.name}: no hay cabecera válida en {titulo}.')
                    finally:
                        libro.close()
                    nueva_firma = ruta.stat()
                    if (firma.st_size, firma.st_mtime_ns) != (nueva_firma.st_size, nueva_firma.st_mtime_ns):
                        raise ValueError(f'{ruta.name} cambió durante la carga. Guarda y vuelve a actualizar.')
                conexion.executemany('INSERT INTO metadatos VALUES (?, ?)',
                                     [('version', VERSION), ('actualizado', datetime.now().astimezone().isoformat(timespec='seconds'))])
                if catalogo is not None:
                    conexion.execute('INSERT INTO metadatos VALUES (?, ?)', ('catalogo_zonas', json.dumps(catalogo, ensure_ascii=False)))
                for ruta, firma in firmas.items():
                    estado = ruta.stat()
                    if (estado.st_size, estado.st_mtime_ns) != firma:
                        raise ValueError(f'{ruta.name} cambió durante la carga. Guarda y vuelve a actualizar.')
            if catalogo is not None:
                catalogo_destino.parent.mkdir(parents=True, exist_ok=True)
                descriptor, nombre = tempfile.mkstemp(dir=catalogo_destino.parent, suffix='.tmp')
                os.close(descriptor)
                catalogo_temporal = Path(nombre)
                catalogo_temporal.write_text(json.dumps(catalogo, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
            if publicacion is not None:
                from .publicacion import publicar_consulta
                publicacion = Path(publicacion)
                publicacion.parent.mkdir(parents=True, exist_ok=True)
                descriptor, nombre = tempfile.mkstemp(dir=publicacion.parent, suffix='.tmp')
                os.close(descriptor)
                consulta_temporal = Path(nombre)
                publicar_consulta(temporal, consulta_temporal)
            reemplazos = [(temporal, destino)]
            if catalogo_temporal is not None: reemplazos.append((catalogo_temporal, catalogo_destino))
            if consulta_temporal is not None: reemplazos.append((consulta_temporal, publicacion))
            anteriores = {p: p.read_bytes() if p.exists() else None for _, p in reemplazos}
            completados = []
            try:
                for nuevo, archivo in reemplazos:
                    os.replace(nuevo, archivo)
                    completados.append((nuevo, archivo))
            except OSError:
                for nuevo, archivo in reversed(completados):
                    if anteriores[archivo] is None: archivo.unlink()
                    else:
                        nuevo.write_bytes(anteriores[archivo])
                        os.replace(nuevo, archivo)
                raise
        finally:
            temporal.unlink(missing_ok=True)
            if consulta_temporal is not None:
                consulta_temporal.unlink(missing_ok=True)
            if catalogo_temporal is not None:
                catalogo_temporal.unlink(missing_ok=True)
        return estado_base_datos(destino)


def leer_configuracion(ruta):
    with closing(conectar(ruta)) as conexion:
        if not conexion.execute("SELECT 1 FROM sqlite_master WHERE name='configuraciones'").fetchone():
            return None
        fila = conexion.execute("SELECT datos FROM configuraciones WHERE clave='portal'").fetchone()
        return json.loads(fila[0]) if fila else None


def guardar_configuracion(ruta, datos, publicacion=None):
    campos = ('regiones', 'cecos', 'catalogo', 'centros', 'ocultos')
    permitidos = set(campos) | {'calendario_excluidas', 'zonas', 'catalogo_zonas', 'unidades_regiones', 'unidades_zonas'}
    if not isinstance(datos, dict) or not set(campos) <= set(datos) or not set(datos) <= permitidos or any(
            not isinstance(valores, list) or any(not isinstance(v, str) for v in valores) for valores in datos.values()):
        raise ValueError('Configuración inválida.')
    # No crea una BD vacía si todavía no se han importado las fuentes.
    estado_base_datos(ruta)
    with ACTUALIZACION, closing(sqlite3.connect(ruta)) as conexion, conexion:
        conexion.execute('CREATE TABLE IF NOT EXISTS configuraciones (clave TEXT PRIMARY KEY, datos TEXT NOT NULL)')
        conexion.execute('INSERT OR REPLACE INTO configuraciones VALUES (?, ?)',
                         ('portal', json.dumps(datos, ensure_ascii=False)))
        if publicacion is not None:
            from .publicacion import actualizar_configuracion_publicada
            actualizar_configuracion_publicada(publicacion, datos, ruta)
    return datos


def conectar(ruta):
    ruta = Path(ruta).resolve()
    if not ruta.is_file():
        raise ValueError('No existe la BD. Pulsa Actualizar BD o ejecuta Actualizar_BD.bat.')
    return sqlite3.connect(ruta.as_uri() + '?mode=ro', uri=True)


def estado_base_datos(ruta):
    with closing(conectar(ruta)) as conexion:
        return _estado(conexion)


def _estado(conexion):
    meta = dict(conexion.execute('SELECT clave, valor FROM metadatos'))
    if meta.get('version') != VERSION:
        raise ValueError('BD incompatible. Ejecuta Actualizar_BD.bat.')
    return dict(actualizado=meta['actualizado'],
                hojas=[dict(fuente=rol, hoja=hoja, filas=max(0, cantidad - 1))
                       for rol, hoja, cantidad in conexion.execute(
                           'SELECT rol, hoja, COUNT(*) FROM filas GROUP BY rol, hoja')])


def leer_fuentes(ruta):
    """Una sola conexión obtiene la misma versión para Excel y maestro."""
    with closing(conectar(ruta)) as conexion, conexion:
        estado = _estado(conexion)
        fuentes = {}
        for rol, nombre, epoch in conexion.execute('SELECT rol, nombre, epoch FROM fuentes'):
            hojas = []
            for (titulo,) in conexion.execute('SELECT nombre FROM hojas WHERE rol=?', (rol,)):
                filas = [tuple(json.loads(datos, object_hook=decodificar)) for (datos,) in conexion.execute(
                    'SELECT datos FROM filas WHERE rol=? AND hoja=? ORDER BY numero', (rol, titulo))]
                hojas.append(HojaBD(titulo, filas))
            fuentes[rol] = FuenteBD(nombre, datetime.fromisoformat(epoch), hojas)
        catalogo = conexion.execute("SELECT valor FROM metadatos WHERE clave='catalogo_zonas'").fetchone()
        if catalogo: fuentes['maestro'].catalogo_zonas = json.loads(catalogo[0])
        return fuentes['excel'], fuentes['maestro'], estado
