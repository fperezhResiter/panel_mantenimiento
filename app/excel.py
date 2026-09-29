"""Lectura de Excel y normalización compartida."""
import unicodedata
from datetime import date, datetime
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

TIPO_ADC = 'Registro Semanal de Kilometrajes y Horómetros (ADC)'
TIPO_MANTENCION = 'Actualización de Mantenciones (Jefe de Mantenimiento)'

def normalizar(valor):
    texto = ' '.join(str(valor or '').split()).upper()
    return ''.join(c for c in unicodedata.normalize('NFD', texto) if not unicodedata.combining(c))


def leer_tabla(ruta, columnas):
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        for hoja in libro:
            filas = hoja.iter_rows(values_only=True)
            cabecera = next(filas, ())
            indices = {normalizar(c): i for i, c in enumerate(cabecera) if c}
            if all(normalizar(c) in indices for c in columnas):
                seleccion = [indices[normalizar(c)] for c in columnas]
                datos = [tuple(f[i] for i in seleccion) for f in filas if any(v is not None for v in f)]
                return datos, libro.epoch, hoja.title
        raise ValueError(f'{ruta.name}: no se encontraron las columnas {", ".join(columnas)} en la primera fila de una hoja.')
    finally:
        libro.close()


def convertir_fecha(valor, epoch):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (float, int)) and not isinstance(valor, bool):
        try:
            return from_excel(valor, epoch).date()
        except (ValueError, OverflowError, AttributeError):
            return None
    if isinstance(valor, str):
        try:
            return datetime.fromisoformat(valor.strip()).date()
        except ValueError:
            for formato in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%d-%m-%Y %H:%M:%S', '%d-%m-%Y %H:%M', '%d-%m-%Y'):
                try:
                    return datetime.strptime(valor.strip(), formato).date()
                except ValueError:
                    pass
    return None


def patente_valida(valor):
    texto = normalizar(valor)
    if texto in ('', '-', 'REVISAR', 'N/A', 'NA', 'SIN PATENTE', 'NO APLICA'):
        return None
    return ''.join(c for c in texto if c not in ' -.') or None


