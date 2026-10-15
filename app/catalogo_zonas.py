"""Lectura del catálogo GDM sin modificar el libro de origen."""
from pathlib import Path
from openpyxl import load_workbook
from .excel import normalizar


def leer_catalogo_zonas(fuente, anterior=None):
    libro = load_workbook(fuente, read_only=True, data_only=True)
    unidades = {}
    try:
        hoja = libro['Prog Resiter']
        for numero, fila in enumerate(hoja.iter_rows(values_only=True), 1):
            if len(fila) < 7:
                continue
            nombre, ceco, ln, linea, zona = fila[2:7]
            if not isinstance(ceco, (int, float)) or not isinstance(ln, (int, float)) or ln >= 50 or not zona:
                continue
            if not nombre or not linea or not float(ceco).is_integer() or not float(ln).is_integer():
                raise ValueError(f'Catálogo GDM: fila {numero} incompleta o con códigos inválidos.')
            clave = (str(nombre).strip(), str(int(ceco)), str(zona).strip())
            u = unidades.setdefault(clave, dict(unidad=clave[0], ceco=clave[1], zona='ZONA '+clave[2], lineas=[]))
            l = dict(codigo=str(int(ln)), nombre=str(linea).strip(), filas_excel=[numero])
            previo = next((x for x in u['lineas'] if x['codigo']==l['codigo'] and x['nombre']==l['nombre']), None)
            if previo: previo['filas_excel'].append(numero)
            else: u['lineas'].append(l)
    finally:
        libro.close()
    if not unidades:
        raise ValueError('No se encontraron asociaciones de unidades y zonas en el catálogo GDM.')
    # Conserva las equivalencias confirmadas; zona y CeCo siempre vienen del Excel.
    for u in unidades.values():
        previos = [p for p in (anterior or {}).get('unidades', []) if normalizar(p['unidad'])==normalizar(u['unidad'])]
        anterior_unidad = next((p for p in previos if p['ceco']==u['ceco']), None)
        if anterior_unidad is None and len(previos)==1:
            anterior_unidad = previos[0]
        if anterior_unidad:
            for campo in ['alias', 'nota']:
                if campo in anterior_unidad: u[campo] = anterior_unidad[campo]
            for l in u['lineas']:
                manual = next((p for p in anterior_unidad['lineas'] if p['codigo']==l['codigo'] and 'nombre_fuente' in p), None)
                if manual:
                    l['nombre_fuente'], l['nombre'] = l['nombre'], manual['nombre']
    return dict(version=1, fuente=Path(fuente).name, hoja='Prog Resiter', unidades=list(unidades.values()))
